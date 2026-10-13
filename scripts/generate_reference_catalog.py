import argparse
import hashlib
import json
import re
import unicodedata
from pathlib import Path

import openpyxl


RULESET_SLUG = "phb-2024-workbook-v2"
LICENSE = "Conteúdo fornecido pelo usuário; licença de redistribuição não informada"


def clean(value):
    return str(value).strip() if value is not None else ""


def slugify(value):
    value = re.sub(r"\s*\([^)]*\)\s*", " ", value)
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def source_info(path):
    data = path.read_bytes()
    return path.name, hashlib.sha256(data).hexdigest()


def rows(sheet, start_row):
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=start_row, max_row=sheet.max_row, values_only=True), start_row
    ):
        values = [clean(value) for value in row]
        if any(values):
            yield row_number, values


def metadata(source_file, source_sha256, tags, **extra):
    return {
        **extra,
        "tags": tags,
        "source_file": source_file,
        "source_sha256": source_sha256,
        "automation": "reference_only",
    }


def weapon_entries(path):
    source_file, source_sha256 = source_info(path)
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    entries = []

    sheet = workbook["Lista de Armas"]
    for row_number, values in rows(sheet, 8):
        name, category, attack_type, damage, damage_type, properties, mastery, weight, cost = values
        entries.append(
            {
                "kind": "item",
                "slug": f"weapon-{slugify(name)}",
                "name": name,
                "summary": f"Arma {category.lower()} de {attack_type.lower()}: {damage} de dano {damage_type.lower()}.",
                "content": "\n".join(
                    [
                        f"Categoria: {category}",
                        f"Alcance/Tipo: {attack_type}",
                        f"Dano: {damage} {damage_type}",
                        f"Propriedades: {properties}",
                        f"Maestria: {mastery}",
                        f"Peso: {weight}",
                        f"Custo: {cost}",
                    ]
                ),
                "source_page": f"Lista de Armas!A{row_number}:I{row_number}",
                "metadata": metadata(
                    source_file,
                    source_sha256,
                    ["Arma", category, attack_type, damage_type],
                    item_type="weapon",
                    category=category,
                    attack_type=attack_type,
                    damage=damage,
                    damage_type=damage_type,
                    properties=properties,
                    weapon_mastery=mastery,
                    weight=weight,
                    cost=cost,
                ),
            }
        )

    sheet = workbook["Propriedades de Maestria"]
    for row_number, values in rows(sheet, 5):
        name, effect, examples = values
        entries.append(
            {
                "kind": "rule",
                "slug": f"weapon-mastery-{slugify(name)}",
                "name": f"Maestria: {name}",
                "summary": effect,
                "content": f"Efeito: {effect}\nExemplos de armas: {examples}",
                "source_page": f"Propriedades de Maestria!A{row_number}:C{row_number}",
                "metadata": metadata(
                    source_file,
                    source_sha256,
                    ["Arma", "Maestria em Arma"],
                    rule_type="weapon_mastery",
                    examples=examples,
                ),
            }
        )

    sheet = workbook["Propriedades Gerais"]
    for row_number, values in rows(sheet, 5):
        name, rule = values
        entries.append(
            {
                "kind": "rule",
                "slug": f"weapon-property-{slugify(name)}",
                "name": f"Propriedade de arma: {name}",
                "summary": rule,
                "content": rule,
                "source_page": f"Propriedades Gerais!A{row_number}:B{row_number}",
                "metadata": metadata(
                    source_file,
                    source_sha256,
                    ["Arma", "Propriedade de arma"],
                    rule_type="weapon_property",
                ),
            }
        )
    return entries


def rules_entries(path):
    source_file, source_sha256 = source_info(path)
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    entries = []

    sheet = workbook["Condições"]
    for row_number, values in rows(sheet, 8):
        name, summary, effects, movement, ending = values
        entries.append(
            {
                "kind": "condition",
                "slug": f"condition-{slugify(name)}",
                "name": name,
                "summary": summary,
                "content": "\n\n".join(
                    [
                        f"Efeitos mecânicos:\n{effects}",
                        f"Movimento e ações:\n{movement}",
                        f"Como encerrar ou suprimir:\n{ending}",
                    ]
                ),
                "source_page": f"Condições!A{row_number}:E{row_number}",
                "metadata": metadata(
                    source_file,
                    source_sha256,
                    ["Condição"],
                    mechanical_effects=effects,
                    movement_impact=movement,
                    ending=ending,
                ),
            }
        )

    sheet = workbook["Talentos"]
    for row_number, values in rows(sheet, 8):
        name, category, prerequisite, ability_score_increase, benefits = values
        first_benefit = next(
            (line.lstrip("•- ").strip() for line in benefits.splitlines() if line.strip()), category
        )
        entries.append(
            {
                "kind": "feat",
                "slug": f"feat-{slugify(name)}",
                "name": name,
                "summary": first_benefit,
                "content": "\n\n".join(
                    [
                        f"Categoria: {category}",
                        f"Pré-requisito: {prerequisite}",
                        f"Aumento de atributo: {ability_score_increase}",
                        f"Benefícios e efeitos:\n{benefits}",
                    ]
                ),
                "source_page": f"Talentos!A{row_number}:E{row_number}",
                "metadata": metadata(
                    source_file,
                    source_sha256,
                    ["Talento", category],
                    category=category,
                    prerequisite=prerequisite,
                    ability_score_increase=ability_score_increase,
                ),
            }
        )
    return entries


def validate(entries):
    expected = {"item": 38, "rule": 18, "condition": 14, "feat": 40}
    actual = {kind: sum(entry["kind"] == kind for entry in entries) for kind in expected}
    if actual != expected:
        raise ValueError(f"Contagens inesperadas: {actual}")
    keys = [(entry["kind"], entry["slug"]) for entry in entries]
    if len(keys) != len(set(keys)):
        raise ValueError("Há chaves kind/slug duplicadas")
    required = ("name", "summary", "content", "source_page", "metadata")
    if any(not entry.get(field) for entry in entries for field in required):
        raise ValueError("Há campos obrigatórios vazios")


def sql_payload(entries):
    payload = json.dumps(entries, ensure_ascii=False, separators=(",", ":"))
    source_files = sorted({entry["metadata"]["source_file"] for entry in entries})
    source_list = ",".join("'" + value.replace("'", "''") + "'" for value in source_files)
    return f"""-- Source workbook import; reference data only. Repeatable without duplicate entries.
begin;
do $import$
declare
 payload jsonb := $payload${payload}$payload$::jsonb;
 row_data jsonb; rid uuid;
begin
 perform pg_advisory_xact_lock(hashtext('grimorio-reference-catalog-2024-import'));
 select id into strict rid from public.dnd_rulesets where slug='{RULESET_SLUG}';
 for row_data in select value from jsonb_array_elements(payload) loop
  insert into public.dnd_rule_entries(ruleset_id,kind,slug,name,summary,content,source_name,source_page,license,metadata)
  values(rid,row_data->>'kind',row_data->>'slug',row_data->>'name',row_data->>'summary',row_data->>'content',
   'Planilhas fornecidas pelo usuário — armas, condições e talentos D&D 2024',row_data->>'source_page',
   '{LICENSE}',row_data->'metadata')
  on conflict(ruleset_id,kind,slug) do update set name=excluded.name,summary=excluded.summary,content=excluded.content,
   source_name=excluded.source_name,source_page=excluded.source_page,license=excluded.license,metadata=excluded.metadata;
 end loop;
 if (select count(*) from public.dnd_rule_entries where metadata->>'source_file' in ({source_list}))<>{len(entries)}
 then raise exception 'Contagem divergente: importação revertida'; end if;
 if exists(select 1 from jsonb_array_elements(payload) p left join public.dnd_rule_entries e
   on e.ruleset_id=rid and e.kind=p->>'kind' and e.slug=p->>'slug'
   where e.id is null or e.name is distinct from p->>'name' or e.summary is distinct from p->>'summary'
    or e.content is distinct from p->>'content' or e.metadata is distinct from p->'metadata')
 then raise exception 'Conteúdo divergente da fonte: importação revertida'; end if;
end;
$import$;
commit;
"""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--weapons", required=True, type=Path)
    parser.add_argument("--rules", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    args = parser.parse_args()
    entries = weapon_entries(args.weapons) + rules_entries(args.rules)
    validate(entries)
    args.out_dir.mkdir(parents=True, exist_ok=True)
    (args.out_dir / "reference_catalog_2024.json").write_text(
        json.dumps(entries, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (args.out_dir / "reference_catalog_2024.sql").write_text(sql_payload(entries), encoding="utf-8")
    counts = {kind: sum(entry["kind"] == kind for entry in entries) for kind in ("item", "rule", "condition", "feat")}
    (args.out_dir / "reference_catalog_2024.md").write_text(
        "# Catálogo de referência D&D 2024\n\n"
        "Carga idempotente criada a partir das duas planilhas fornecidas pelo usuário.\n\n"
        f"- Armas: {counts['item']}\n"
        f"- Regras de propriedades e maestrias de armas: {counts['rule']}\n"
        f"- Condições: {counts['condition']}\n"
        f"- Talentos: {counts['feat']}\n"
        f"- Total: {len(entries)}\n\n"
        "Os registros são referências para consulta no glossário e não automatizam alterações na ficha.\n",
        encoding="utf-8",
    )
    print(json.dumps({"total": len(entries), "counts": counts}, ensure_ascii=False))


if __name__ == "__main__":
    main()
